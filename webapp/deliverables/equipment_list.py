"""Equipment List deliverable generators (CSV + XLSX)."""

import csv
import glob
import io as _io
import json
from pathlib import Path
from typing import ClassVar, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from webapp.deliverables.base import Generator
from webapp.deliverables.canonical import CanonicalEntity, JobCanonical
from webapp.deliverables.field_resolver import resolve_field
from webapp.deliverables.registry import REGISTRY
from webapp.deliverables.sort_utils import instrument_sort_key
from webapp.deliverables.template import DeliverableConfig, TemplateConfig
from webapp.deliverables.xlsx_utils import autosize_columns, fit_to_one_page

# Maps equipment_specs.json keys → CanonicalEntity.fields keys
_SPEC_FIELD_MAP = {
    "type":                  "equip_spec_type",
    "rated_capacity":        "rated_capacity",
    "differential_pressure": "differential_pressure",
    "design_temperature":    "design_temp",
    "motor_rating":          "motor_rating",
    "material":              "material",
    "quantity":              "quantity",
}


def _load_equipment_specs(job_id: int) -> dict:
    """Load equipment_specs.json for job_id; returns {} if absent or unreadable."""
    from webapp.config import JOB_OUTPUT_DIR
    candidates = sorted(glob.glob(str(JOB_OUTPUT_DIR / "*" / str(job_id) / "equipment_specs.json")))
    if candidates:
        try:
            return json.loads(Path(candidates[0]).read_text()) or {}
        except Exception:
            pass
    legacy = JOB_OUTPUT_DIR / str(job_id) / "equipment_specs.json"
    if legacy.exists():
        try:
            return json.loads(legacy.read_text()) or {}
        except Exception:
            pass
    return {}


def _enrich_equipment(entities: List[CanonicalEntity], specs: dict) -> None:
    """Inject Pass 5 spec fields into entity.fields in-place (only fills blanks)."""
    for e in entities:
        tag = e.tag or ""
        spec = specs.get(tag) or {}
        if not spec:
            continue
        for spec_key, field_key in _SPEC_FIELD_MAP.items():
            val = spec.get(spec_key)
            if val is None:
                continue
            if isinstance(val, str):
                val = val.strip()
            if val and str(val).lower() not in ("null", "none"):
                e.fields.setdefault(field_key, val)


def _filter_equipment(entities: List[CanonicalEntity]) -> List[CanonicalEntity]:
    return [e for e in entities if e.entity_class == "equipment"]


def _ordered_columns(cfg: DeliverableConfig):
    return sorted(cfg.columns, key=lambda c: c.order)


class EquipmentListCSVGenerator(Generator):
    deliverable_type: ClassVar[str] = "equipment_list"
    file_format: ClassVar[str] = "csv"

    def generate(self, canonical: JobCanonical, template: TemplateConfig) -> bytes:
        deliv = template.deliverables["equipment_list"]
        columns = _ordered_columns(deliv)
        equipment = _filter_equipment(canonical.entities)
        _enrich_equipment(equipment, _load_equipment_specs(canonical.job_id))
        equipment.sort(key=lambda e: instrument_sort_key(e.tag or ""))
        buf = _io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([c.header for c in columns])
        for e in equipment:
            writer.writerow([resolve_field(e, c.field) for c in columns])
        return buf.getvalue().encode("utf-8")


class EquipmentListXLSXGenerator(Generator):
    deliverable_type: ClassVar[str] = "equipment_list"
    file_format: ClassVar[str] = "xlsx"

    def generate(self, canonical: JobCanonical, template: TemplateConfig) -> bytes:
        deliv = template.deliverables["equipment_list"]
        columns = _ordered_columns(deliv)
        equipment = _filter_equipment(canonical.entities)
        _enrich_equipment(equipment, _load_equipment_specs(canonical.job_id))
        equipment.sort(key=lambda e: instrument_sort_key(e.tag or ""))
        wb = Workbook()
        ws = wb.active
        ws.title = deliv.sheet_name or "Equipment List"
        ws.append([c.header for c in columns])
        header_font = Font(name=deliv.font, bold=True)
        header_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        for e in equipment:
            ws.append([resolve_field(e, c.field) for c in columns])
        # Fit the whole list onto a single page (adapts to A4/A3 at print time).
        autosize_columns(ws)
        fit_to_one_page(ws)
        buf = _io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


REGISTRY.register(EquipmentListCSVGenerator)
REGISTRY.register(EquipmentListXLSXGenerator)
