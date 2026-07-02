"""Shared XLSX helpers for deliverable exports.

`fit_to_one_page` configures an openpyxl worksheet so its whole table prints /
exports to a single page via Excel's Page Setup "Fit to page" scaling. The fit
is paper-size-agnostic: Excel rescales to one page of whatever paper the user
selects at print time (A4, A3, …), so the same workbook works for both.
"""
from __future__ import annotations

from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties


def autosize_columns(ws: Any, *, min_width: int = 8, max_width: int = 60) -> None:
    """Set each column's width from its longest cell so nothing is clipped
    before the fit-to-page scaling is applied."""
    widths: dict = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col = cell.column  # 1-based int
            length = len(str(cell.value))
            if length > widths.get(col, 0):
                widths[col] = length
    for col, length in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(
            min_width, min(max_width, length + 2)
        )


def fit_to_one_page(ws: Any, *, landscape: bool = True) -> None:
    """Make the worksheet print/export onto exactly ONE page.

    Sets Page Setup scaling to 1 page wide x 1 page tall and enables the
    fit-to-page flag (without it Excel ignores fitToWidth/fitToHeight). Paper
    size is intentionally left at the workbook default so the single-page fit
    adapts to whichever paper (A4/A3/…) is chosen when printing.
    """
    if landscape:
        ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    # Required — Excel honours fitToWidth/fitToHeight only when this flag is on.
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    # Narrow margins give the scaler more usable area before shrinking text.
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.25
