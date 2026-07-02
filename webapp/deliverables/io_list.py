"""I/O List deliverable generators (CSV + XLSX).

Shows all instruments (entity_class == "instrument") with their DCS/PLC
signal and wiring data. The io_type column (AI/AO/DI/DO) is empty for
locally-mounted instruments that have no control-system connection.
"""

import csv
import io as _io
from typing import ClassVar, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from webapp.deliverables.base import Generator
from webapp.deliverables.canonical import CanonicalEntity, JobCanonical
from webapp.deliverables.field_resolver import resolve_field
from webapp.deliverables.registry import REGISTRY
from webapp.deliverables.sort_utils import instrument_sort_key
from webapp.deliverables.template import DeliverableConfig, TemplateConfig
from webapp.deliverables.xlsx_utils import autosize_columns, fit_to_one_page


# Excel-ONLY grouped headers for the IO List (the UI/template are untouched —
# this mapping lives in the generator). Maps the flat template header → its
# (group label, short sub-header). Grouped columns render as a merged group
# cell in header row 1 with the short label in row 2; everything else is a
# single header merged across both rows.
_XLSX_HEADER_GROUPS = {
    "Junction TB-1":                ("Junction TB No.", "TB-1"),
    "Junction TB-2":                ("Junction TB No.", "TB-2"),
    "Marshalling Cabinet TB Strip": ("Marshalling Cabinet Terminal Block No.", "TB Strip"),
    "Marshalling Cabinet TB-1":     ("Marshalling Cabinet Terminal Block No.", "TB-1"),
    "Marshalling Cabinet TB-2":     ("Marshalling Cabinet Terminal Block No.", "TB-2"),
}


def _filter_instruments(entities: List[CanonicalEntity]) -> List[CanonicalEntity]:
    return [e for e in entities if e.entity_class == "instrument"]


def _ordered_columns(cfg: DeliverableConfig):
    return sorted(cfg.columns, key=lambda c: c.order)


class IOListCSVGenerator(Generator):
    deliverable_type: ClassVar[str] = "io_list"
    file_format: ClassVar[str] = "csv"

    def generate(self, canonical: JobCanonical, template: TemplateConfig) -> bytes:
        deliv = template.deliverables["io_list"]
        columns = _ordered_columns(deliv)
        instruments = _filter_instruments(canonical.entities)
        instruments.sort(key=lambda e: instrument_sort_key(e.tag or ""))
        buf = _io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([c.header for c in columns])
        for inst in instruments:
            writer.writerow([resolve_field(inst, c.field) for c in columns])
        return buf.getvalue().encode("utf-8")


class IOListXLSXGenerator(Generator):
    deliverable_type: ClassVar[str] = "io_list"
    file_format: ClassVar[str] = "xlsx"

    def generate(self, canonical: JobCanonical, template: TemplateConfig) -> bytes:
        deliv = template.deliverables["io_list"]
        columns = _ordered_columns(deliv)
        instruments = _filter_instruments(canonical.entities)
        instruments.sort(key=lambda e: instrument_sort_key(e.tag or ""))
        wb = Workbook()
        ws = wb.active
        raw_title = deliv.sheet_name or "IO List"
        ws.title = raw_title.replace("/", "-").replace("\\", "-")[:31]

        header_font = Font(name=deliv.font, bold=True)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        # All header cells (including the TB group cells) use the same yellow.
        yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        def _style(cell, fill):
            cell.font = header_font
            cell.alignment = center
            cell.fill = fill
            cell.border = border

        # ── Two-row header: group (row 1, merged) + sub-header (row 2) ───────
        # Per column, resolve (group_label or None, label). Ungrouped columns
        # use the full header and merge vertically across both header rows.
        meta = [(_XLSX_HEADER_GROUPS.get(c.header, (None, c.header))) for c in columns]

        for i, (group_label, label) in enumerate(meta, start=1):
            col = get_column_letter(i)
            if group_label is None:
                c1 = ws.cell(row=1, column=i, value=label)
                _style(c1, yellow)
                _style(ws.cell(row=2, column=i), yellow)
                ws.merge_cells(f"{col}1:{col}2")
            else:
                _style(ws.cell(row=2, column=i, value=label), yellow)

        # Merge each contiguous run of same-group columns across header row 1.
        start = 0
        while start < len(meta):
            grp = meta[start][0]
            if grp is None:
                start += 1
                continue
            end = start
            while end + 1 < len(meta) and meta[end + 1][0] == grp:
                end += 1
            first, last = start + 1, end + 1
            for ci in range(first, last + 1):
                _style(ws.cell(row=1, column=ci), yellow)
            ws.cell(row=1, column=first, value=grp)
            if first != last:
                ws.merge_cells(f"{get_column_letter(first)}1:{get_column_letter(last)}1")
            start = end + 1

        # ── Data rows (row 3 onwards) ────────────────────────────────────────
        for inst in instruments:
            ws.append([resolve_field(inst, c.field) for c in columns])

        # Fit the whole list onto a single page (adapts to A4/A3 at print time).
        autosize_columns(ws)
        fit_to_one_page(ws)
        buf = _io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


REGISTRY.register(IOListCSVGenerator)
REGISTRY.register(IOListXLSXGenerator)
