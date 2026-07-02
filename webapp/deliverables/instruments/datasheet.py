"""Instrument Datasheet XLSX Export — engineering specification-sheet layout.

Reproduces the visual structure of a real process-instrument specification
sheet (the TechnipEnergies/Reliance "INSTRUMENT SPECIFICATION" form):

    * a grey classification banner across the top
    * a full-width upper zone (identity / process / certification ...),
      laid out as   section-band | # | label | value
    * a two-column lower zone packing the detailed type-specific sections
      side by side (left panel + right panel), each   band | # | label | value
    * a "Note:" block
    * an engineering title block at the bottom (MOC No. / JOB ID / Program /
      DEC-PV, the INSTRUMENT SPECIFICATION + type strip with a logo cell,
      Code / Sheet / Rev, and a No|By|Chk|Appr|Date|Revision|DS.No. table)

Section bands and header strips carry a light-grey fill; every cell has a thin
black border; each sheet prints A4-landscape fit-to-page.

The CONTENT (the labels) is driven entirely by the IDS schema — the exact same
fields the Bulk Review right-side panel shows (``get_ids_sections_for_type``).
Values are written when present (overrides / vendor data) and left blank
otherwise.  Every field appears once; no field is repeated.
"""

import io
import logging
from typing import ClassVar, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from webapp.deliverables.base import Generator
from webapp.deliverables.canonical import CanonicalEntity, JobCanonical
from webapp.deliverables.ids_schema import (
    TYPE_LABELS,
    get_ids_sections_for_type,
    normalize_subclass,
)
from webapp.deliverables.registry import REGISTRY
from webapp.deliverables.template import TemplateConfig

log = logging.getLogger(__name__)

_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_GREY = PatternFill("solid", fgColor="D9D9D9")

# 10-column grid (proportions mirror the reference form).
#   Upper (full-width):  A band | B # | C:E label | F:J value
#   Lower left panel:    A band | B # | C:D label | E:F value
#   Lower right panel:   G band | H # | I label   | J value
_NCOLS = 10
_COL_WIDTHS = {
    "A": 15, "B": 4, "C": 16, "D": 10, "E": 30,
    "F": 14, "G": 15, "H": 4, "I": 26, "J": 30,
}
_ROW_H = 14.0

# (band, num, label-span, value-span) column layouts.
_UPPER = (1, 2, (3, 5), (6, 10))
_LOWER_L = (1, 2, (3, 4), (5, 6))
_LOWER_R = (7, 8, (9, 9), (10, 10))


# Vendor-dependent fields stored under short keys in entity.fields.
# Bridged to IDS schema paths so exports populate correctly.
_VENDOR_FIELD_ALIASES: Dict[str, str] = {
    "fields.piping_class":        "fields.ids_pipe_class",
    "fields.calb_range_min":      "fields.ids_calibration_range_min",
    "fields.calb_range_max":      "fields.ids_calibration_range_max",
    "fields.calb_range_unit":     "fields.ids_calibration_range_unit",
    "fields.measuring_range_min": "fields.ids_instrument_range_min",
    "fields.measuring_range_max": "fields.ids_instrument_range_max",
    "fields.measuring_range_unit":"fields.ids_instrument_range_unit",
    "fields.certification":       "fields.ids_certification_special_requirement",
}

# Vendor extras not in IDS schema — appended as labelled rows when they carry a value.
_VENDOR_EXTRAS: List[Tuple[str, str]] = [
    ("Power Supply Input",  "fields.power_in"),
    ("Power Supply Output", "fields.power_out"),
    ("Output Signal",       "fields.io_output"),
    ("Datasheet Ref",       "fields.datasheet_ref"),
]


def _safe_name(name: str) -> str:
    safe = "".join(c if c not in r"[]:*?/\\" else "_" for c in name)
    return safe[:31]


def _entity_to_value_dict(entity: CanonicalEntity) -> Dict[str, str]:
    """Resolve all entity field values in a single pass → {path: str_value}."""
    data: Dict[str, str] = {
        "tag":        entity.tag or "",
        "pid_number": entity.pid_number or "",
        "sub_class":  entity.sub_class or "",
    }
    for k, v in (entity.fields or {}).items():
        data[f"fields.{k}"] = str(v) if v is not None else ""
    for short_path, ids_path in _VENDOR_FIELD_ALIASES.items():
        if data.get(short_path):
            data.setdefault(ids_path, data[short_path])
    if entity.vendor_match:
        vm = entity.vendor_match
        data["vendor_match.vendor_name"]  = vm.vendor_name or ""
        data["vendor_match.product_name"] = vm.product_name or ""
        data["vendor_match.part_number"]  = vm.part_number or ""
        for k, v in (vm.catalog_fields or {}).items():
            data[f"vendor_match.catalog_fields.{k}"] = str(v) if v is not None else ""
    return data


# A "block" = (SECTION NAME, [(label, value), ...]).
Block = Tuple[str, List[Tuple[str, str]]]


def _build_sections(entity: CanonicalEntity,
                    value_dict: Dict[str, str]) -> List[Block]:
    """Deduped section blocks for one entity, in panel order.

    Field order/grouping mirror the Bulk Review right-side panel exactly.
    A field path is emitted at most once (no repeats); empty sections drop out.
    """
    seen: set = set()
    built: List[Block] = []
    for section in get_ids_sections_for_type(entity.sub_class):
        rows: List[Tuple[str, str]] = []
        for field in section.fields:
            if field.path in seen:
                continue
            seen.add(field.path)
            rows.append((field.header, value_dict.get(field.path, "")))
        if rows:
            built.append((section.name.upper(), rows))

    extras = [(label, value_dict[path])
              for label, path in _VENDOR_EXTRAS
              if path not in seen and value_dict.get(path)]
    if extras:
        built.append(("VENDOR DATA", extras))
    return built


# The reference form runs fields 1-26 full-width, then 27-82 two-column.
# We mirror that proportion: the full-width upper zone ends at the section
# boundary whose cumulative field count is closest to this target.
_UPPER_TARGET = 26


def _zones(blocks: List[Block]) -> Tuple[List[Block], List[Block]]:
    """Split section blocks into (full-width upper, two-column lower).

    The cut falls on a section boundary nearest ``_UPPER_TARGET`` fields, so
    sections are never broken across zones.
    """
    if not blocks:
        return [], []
    cums, acc = [], 0
    for _, rows in blocks:
        acc += len(rows)
        cums.append(acc)
    best_k = min(range(1, len(blocks) + 1),
                 key=lambda k: abs(cums[k - 1] - _UPPER_TARGET))
    return blocks[:best_k], blocks[best_k:]


def _partition(blocks: List[Block]) -> Tuple[List[Block], List[Block]]:
    """Split lower-zone blocks into two balanced panels at a section edge."""
    total = sum(len(rows) for _, rows in blocks)
    left: List[Block] = []
    acc = 0
    half = total / 2.0
    for i, blk in enumerate(blocks):
        left.append(blk)
        acc += len(blk[1])
        if acc >= half and i < len(blocks) - 1:
            return left, blocks[i + 1:]
    return blocks, []


class DatasheetXLSXGenerator(Generator):
    """One XLSX with one engineering-spec sheet per instrument/valve entity."""

    deliverable_type: ClassVar[str] = "datasheet"
    file_format: ClassVar[str] = "xlsx"

    def generate(self, canonical: JobCanonical, _template: TemplateConfig,
                 active_sheet: Optional[str] = None) -> bytes:
        entities = [
            e for e in canonical.entities
            if e.entity_class in ("instrument", "valve")
        ]
        entities.sort(key=lambda e: e.tag or "")

        wb_out = Workbook()
        if wb_out.active:
            wb_out.remove(wb_out.active)

        if not entities:
            ws = wb_out.create_sheet(title="No Data")
            ws["A1"] = "No instrument or valve entities found in this job."
            out = io.BytesIO()
            wb_out.save(out)
            return out.getvalue()

        for entity in entities:
            self._fill_vendor_data(entity)
            value_dict = _entity_to_value_dict(entity)
            built = _build_sections(entity, value_dict)

            ws = wb_out.create_sheet(title=_safe_name(entity.tag or str(entity.entity_id)))
            self._render_sheet(ws, entity, built)
            log.info("Datasheet sheet: %s (%s) — %d sections",
                     entity.tag, entity.sub_class, len(built))

        if active_sheet:
            safe = _safe_name(active_sheet)
            if safe in wb_out.sheetnames:
                wb_out.active = wb_out[safe]

        out = io.BytesIO()
        wb_out.save(out)
        return out.getvalue()

    # ------------------------------------------------------------------ render
    def _render_sheet(self, ws, entity: CanonicalEntity, built: List[Block]) -> None:
        for letter, width in _COL_WIDTHS.items():
            ws.column_dimensions[letter].width = width

        full, detail = _zones(built)

        # --- top classification banner ----------------------------------
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=_NCOLS)
        self._cell(ws, 1, 1, "INSTRUMENT DATA SHEET", bold=True, size=10, fill=_GREY)
        ws.row_dimensions[1].height = 16
        row, seq = 2, 1

        # --- full-width upper zone --------------------------------------
        row, seq = self._panel(ws, full, row, seq, _UPPER)

        # --- two-column lower zone --------------------------------------
        left, right = _partition(detail)
        _, seq = self._panel(ws, left, row, seq, _LOWER_L)
        self._panel(ws, right, row, seq, _LOWER_R)
        row += max(self._height(left), self._height(right))

        # --- note block -------------------------------------------------
        self._cell(ws, row, 1, "Note:", bold=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=_NCOLS)
        ws.row_dimensions[row].height = _ROW_H
        row += 1
        for _ in range(3):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=_NCOLS)
            ws.row_dimensions[row].height = _ROW_H
            row += 1

        # --- title block (bottom) ---------------------------------------
        last = self._title_block(ws, entity, row)

        # --- continuous border grid -------------------------------------
        for r in range(1, last + 1):
            for c in range(1, _NCOLS + 1):
                ws.cell(row=r, column=c).border = _BORDER

        # --- print setup ------------------------------------------------
        ws.page_setup.paperSize = 9             # A4
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4,
                                      header=0.2, footer=0.2)
        ws.sheet_view.showGridLines = False

    # --- panels -------------------------------------------------------------
    def _height(self, blocks: List[Block]) -> int:
        return sum(len(rows) for _, rows in blocks)

    def _panel(self, ws, blocks, start_row, seq, layout):
        """Render section blocks down one panel; return (next_row, next_seq)."""
        band_c, num_c, (lc0, lc1), (vc0, vc1) = layout
        row = start_row
        for name, rows in blocks:
            band_start = row
            for label, value in rows:
                self._cell(ws, row, num_c, seq, align="center", wrap=False)
                self._span(ws, row, lc0, lc1, label, bold=True)
                self._span(ws, row, vc0, vc1, value)
                ws.row_dimensions[row].height = _ROW_H
                row += 1
                seq += 1
            self._band(ws, band_c, band_start, row - 1, name)
        return row, seq

    def _band(self, ws, col, r0, r1, text):
        if r1 < r0:
            return
        if r1 > r0:
            ws.merge_cells(start_row=r0, start_column=col, end_row=r1, end_column=col)
        self._cell(ws, r0, col, text, bold=True, size=8, align="center", fill=_GREY)

    def _span(self, ws, row, c0, c1, value, **kw):
        self._cell(ws, row, c0, value, **kw)
        if c1 > c0:
            ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c1)

    def _cell(self, ws, row, col, value, *, bold=False, size=8,
              align="left", wrap=True, fill=None):
        c = ws.cell(row=row, column=col, value=value)
        c.font = Font(name="Calibri", bold=bold, size=size)
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if fill is not None:
            c.fill = fill

    # --- bottom engineering title block -------------------------------------
    def _title_block(self, ws, entity: CanonicalEntity, row: int) -> int:
        type_label = TYPE_LABELS.get(normalize_subclass(entity.sub_class) or "",
                                     "Instrument")

        def hdr(r, spans):
            for (c0, c1, text) in spans:
                self._span(ws, r, c0, c1, text, bold=True, align="center", fill=_GREY)
            ws.row_dimensions[r].height = _ROW_H

        def blank(r, spans):
            for (c0, c1) in spans:
                self._span(ws, r, c0, c1, "")
            ws.row_dimensions[r].height = _ROW_H

        # MOC / JOB ID / Program / DEC-PV
        hdr(row, [(1, 5, "MOC No."), (6, 7, "JOB ID"),
                  (8, 8, "Program No."), (9, 10, "DEC / PV")])
        blank(row + 1, [(1, 5), (6, 7), (8, 8), (9, 10)])

        # INSTRUMENT SPECIFICATION + type   |   logo cell
        r = row + 2
        self._span(ws, r, 1, 7, f"INSTRUMENT SPECIFICATION   —   {type_label}",
                   bold=True, size=11, align="center")
        self._span(ws, r, 8, 10, "(Company Logo)", align="center", fill=_GREY)
        ws.row_dimensions[r].height = 30

        # Code / Sheet / of / Rev
        self._span(ws, row + 3, 1, 5, "Code:", bold=True)
        self._span(ws, row + 3, 6, 6, "Sheet", bold=True, align="center")
        self._span(ws, row + 3, 7, 7, "of", bold=True, align="center")
        self._span(ws, row + 3, 8, 10, "Rev.:", bold=True)
        ws.row_dimensions[row + 3].height = _ROW_H

        # Revision table header + two blank entry rows
        hdr(row + 4, [(1, 1, "No."), (2, 2, "By"), (3, 3, "Chk"), (4, 4, "Appr"),
                      (5, 6, "Date"), (7, 8, "Revision"), (9, 10, "DS. No.:")])
        blank(row + 5, [(1, 1), (2, 2), (3, 3), (4, 4), (5, 6), (7, 8), (9, 10)])
        blank(row + 6, [(1, 1), (2, 2), (3, 3), (4, 4), (5, 6), (7, 8), (9, 10)])
        return row + 6

    # -------------------------------------------------------------- vendor data
    def _fill_vendor_data(self, entity: CanonicalEntity) -> None:
        """Auto-fetch vendor datasheet — only fills EMPTY fields."""
        if entity.entity_class not in ("instrument", "valve") or entity.vendor_match:
            return
        from webapp.deliverables.vendor_match_client import fetch_vendor_datasheet
        import uuid

        tag = entity.tag or ""
        # Take prefix before first hyphen: "PT-3069A" → "PT", "PIT-3087" → "PIT".
        # Normalise 3-letter codes: "PIT" → "PT", "FIT" → "FT", "TIT" → "TT".
        inst_code = tag.split("-")[0].upper()
        if len(inst_code) > 2 and inst_code.endswith("T"):
            inst_code = inst_code[0] + "T"
        elif len(inst_code) > 2 and inst_code.endswith("E"):
            inst_code = inst_code[0] + "E"
        if not inst_code or len(inst_code) > 4:
            return
        vendor_data = fetch_vendor_datasheet(inst_code)
        if not vendor_data:
            return
        for k, v in vendor_data.items():
            if k.startswith("ids_") and v and not entity.fields.get(k):
                entity.fields[k] = v
        if vendor_data.get("_vendor_name") or vendor_data.get("_model_number"):
            from webapp.deliverables.canonical import VendorMatch
            catalog = {k: v for k, v in vendor_data.items()
                       if not k.startswith("_") and not k.startswith("ids_") and v}
            entity.vendor_match = VendorMatch(
                vendor_id=str(uuid.uuid5(uuid.NAMESPACE_DNS, str(entity.entity_id))),
                vendor_name=vendor_data.get("_vendor_name", ""),
                product_name=vendor_data.get("_model_number", ""),
                part_number=vendor_data.get("_product_id", ""),
                catalog_fields=catalog,
            )


# Backwards-compatible alias.
ProfessionalDatasheetXLSXGenerator = DatasheetXLSXGenerator

REGISTRY.register(DatasheetXLSXGenerator)
