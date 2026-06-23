import io
import json
import openpyxl
from pathlib import Path

from webapp.deliverables.canonical import JobCanonical
from webapp.deliverables.template_loader import TemplateLoader
from webapp.deliverables.valve_list import ValveListCSVGenerator, ValveListXLSXGenerator

FIXTURE_DIR = Path(__file__).parent / "fixtures"


def load_sample() -> JobCanonical:
    raw = json.loads((FIXTURE_DIR / "canonical_sample.json").read_text())
    return JobCanonical.model_validate(raw)


def test_valvelist_csv_default_template_matches_production_format():
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out = ValveListCSVGenerator().generate(job, tpl)
    expected = (FIXTURE_DIR / "expected_valve_list_default.csv").read_bytes()
    assert out == expected


def test_valvelist_csv_filters_to_valves_only():
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out = ValveListCSVGenerator().generate(job, tpl).decode("utf-8")
    assert "62-GL-151000" in out or "151000" in out
    assert "62-BV-151001" in out or "151001" in out
    assert "422-11-PT-006A" not in out
    assert "V-101" not in out


def test_valvelist_csv_preserves_trailing_space_in_pneumatic_actuator_header():
    """Production header "Pneumatic Actuator " has a trailing space. Some agency
    tooling parses by exact column name. Must be preserved verbatim."""
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out = ValveListCSVGenerator().generate(job, tpl).decode("utf-8")
    header_line = out.splitlines()[0]
    assert "Pneumatic Actuator " in header_line  # note trailing space
    assert "Pneumatic Actuator," not in header_line  # the no-space version must NOT appear


def test_valvelist_csv_class_attrs():
    assert ValveListCSVGenerator.deliverable_type == "valve_list"
    assert ValveListCSVGenerator.file_format == "csv"


EXPECTED_VALVE_LIST_HEADERS = [
    "P&ID No", "Dynamic Code", "Category", "Size", "Area Code",
    "Serial No", "Series Code", "Fluid Code", "Piping Class", "Qty",
    "Motor Actuator", "Pneumatic Actuator ", "Solenoid", "line",  # note trailing space on Pneumatic Actuator
]


def test_valvelist_xlsx_default_template():
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out_bytes = ValveListXLSXGenerator().generate(job, tpl)
    wb = openpyxl.load_workbook(io.BytesIO(out_bytes))
    ws = wb["Valve List"]
    headers = [c.value for c in ws[1]]
    assert headers == EXPECTED_VALVE_LIST_HEADERS
    row2 = [c.value for c in ws[2]]
    assert row2 == [
        "MUK-62-1-15-1001-001-24C7-D", "-", "GL", "8", "62",
        "151000", "-", "G", "AC-PP", "1",
        "-", "-", "-", "8\"-G-62151004-AC-PP",
    ]


def test_valvelist_xlsx_uses_ronesans_sheet_name():
    job = load_sample()
    tpl = TemplateLoader().load("ronesans")
    out_bytes = ValveListXLSXGenerator().generate(job, tpl)
    wb = openpyxl.load_workbook(io.BytesIO(out_bytes))
    assert "VALVE LIST" in wb.sheetnames


def test_valvelist_xlsx_class_attrs():
    assert ValveListXLSXGenerator.deliverable_type == "valve_list"
    assert ValveListXLSXGenerator.file_format == "xlsx"
