import io
import json
from pathlib import Path

import openpyxl

from webapp.deliverables.canonical import JobCanonical
from webapp.deliverables.equipment_list import EquipmentListXLSXGenerator
from webapp.deliverables.template_loader import TemplateLoader

FIXTURE_DIR = Path(__file__).parent / "fixtures"


def load_sample() -> JobCanonical:
    raw = json.loads((FIXTURE_DIR / "canonical_sample.json").read_text())
    return JobCanonical.model_validate(raw)


def test_equipmentlist_xlsx_default():
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out_bytes = EquipmentListXLSXGenerator().generate(job, tpl)
    wb = openpyxl.load_workbook(io.BytesIO(out_bytes))
    ws = wb["Equipment List"]
    headers = [c.value for c in ws[1]]
    assert headers == [
        "Unit Number", "Equipment Tag", "Equipment Type", "Service Description",
        "Quantity", "Type", "Rated Capacity", "Diff. Pressure", "Design Temp.",
        "Motor Rating", "Material", "P&ID No.", "Location", "Manufacturer",
        "Model No", "Remark",
    ]
    row2 = [c.value for c in ws[2]]
    assert row2 == [
        "422-11", "V-101", "Reactor Vessel", None, None, None,
        None, None, None, None, None, "VEN-M5BC-6-50-0006", "FIELD",
        None, None, None,
    ]


def test_equipmentlist_filters_to_equipment():
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out_bytes = EquipmentListXLSXGenerator().generate(job, tpl)
    wb = openpyxl.load_workbook(io.BytesIO(out_bytes))
    ws = wb["Equipment List"]
    # Equipment Tag column is column 2 in default template
    tags = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert "V-101" in tags
    assert "62-GL-151000" not in tags
    assert "422-11-PT-006A" not in tags


def test_equipmentlist_class_attrs():
    assert EquipmentListXLSXGenerator.deliverable_type == "equipment_list"
    assert EquipmentListXLSXGenerator.file_format == "xlsx"
