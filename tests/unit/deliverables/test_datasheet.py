import io
import json
from pathlib import Path

import openpyxl

from webapp.deliverables.canonical import JobCanonical
from webapp.deliverables.datasheet import DatasheetXLSXGenerator, IDS_SECTIONS
from webapp.deliverables.template_loader import TemplateLoader

FIXTURE_DIR = Path(__file__).parent / "fixtures"


def load_sample() -> JobCanonical:
    raw = json.loads((FIXTURE_DIR / "canonical_sample.json").read_text())
    return JobCanonical.model_validate(raw)


def open_workbook(out_bytes: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(out_bytes))


def test_datasheet_xlsx_has_one_sheet_per_instrument():
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out = DatasheetXLSXGenerator().generate(job, tpl)
    wb = open_workbook(out)
    instrument_count = sum(1 for e in job.entities if e.entity_class == "instrument")
    assert len(wb.sheetnames) == instrument_count


def test_datasheet_xlsx_sheet_named_after_tag():
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out = DatasheetXLSXGenerator().generate(job, tpl)
    wb = open_workbook(out)
    assert "422-11-PT-006A" in wb.sheetnames


def test_datasheet_xlsx_has_all_11_sections():
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out = DatasheetXLSXGenerator().generate(job, tpl)
    wb = open_workbook(out)
    ws = wb["422-11-PT-006A"]
    # All section labels should appear in column A
    column_a_values = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    section_names = [name for name, _fields in IDS_SECTIONS]
    assert len(section_names) == 11
    for section in section_names:
        assert section in column_a_values, f"section '{section}' missing"


def test_datasheet_xlsx_contains_vendor_data_in_purchase_section():
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out = DatasheetXLSXGenerator().generate(job, tpl)
    wb = open_workbook(out)
    ws = wb["422-11-PT-006A"]
    all_cell_values = [
        ws.cell(row=r, column=c).value
        for r in range(1, ws.max_row + 1)
        for c in range(1, ws.max_column + 1)
    ]
    assert "Yokogawa" in all_cell_values
    assert "EJX130A" in all_cell_values
    assert "EJX130A-JMS5G-022NN" in all_cell_values


def test_datasheet_xlsx_contains_general_data_fields():
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out = DatasheetXLSXGenerator().generate(job, tpl)
    wb = open_workbook(out)
    ws = wb["422-11-PT-006A"]
    all_cell_values = [
        ws.cell(row=r, column=c).value
        for r in range(1, ws.max_row + 1)
        for c in range(1, ws.max_column + 1)
    ]
    assert "422-11-PT-006A" in all_cell_values
    assert "COMP SUCTION PRESS" in all_cell_values
    assert "VEN-M5BC-6-50-0002" in all_cell_values


def test_datasheet_xlsx_section_headers_are_merged_a_to_c():
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out = DatasheetXLSXGenerator().generate(job, tpl)
    wb = open_workbook(out)
    ws = wb["422-11-PT-006A"]
    merged_ranges = {str(r) for r in ws.merged_cells.ranges}
    # Find the row where "General Data" appears in column A; expect A:C merge there
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "General Data":
            assert f"A{r}:C{r}" in merged_ranges
            return
    raise AssertionError("General Data section header row not found")


def test_datasheet_xlsx_class_attrs():
    assert DatasheetXLSXGenerator.deliverable_type == "datasheet"
    assert DatasheetXLSXGenerator.file_format == "xlsx"


def test_datasheet_xlsx_starts_with_zip_magic():
    """XLSX files are ZIP archives — magic bytes PK\\x03\\x04."""
    job = load_sample()
    tpl = TemplateLoader().load("default")
    out = DatasheetXLSXGenerator().generate(job, tpl)
    assert out[:2] == b"PK"
